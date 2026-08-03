from __future__ import annotations

from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data" / "exports" / "workflow-doc"
DOCX_PATH = OUT_DIR / "ETF主题雷达Agent工作流程与技术架构说明.docx"
FLOW_IMAGE = OUT_DIR / "workflow.png"
STATE_IMAGE = OUT_DIR / "state-machine.png"

NAVY = "17324D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8F1F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "68717A"
GRID = "CCD4DC"
WHITE = "FFFFFF"
GREEN = "287A52"
AMBER = "A66A00"
RED = "9B1C1C"


def set_font(run, size: float | None = None, bold: bool | None = None,
             color: str | None = None, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120,
                     bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header: bool = True) -> None:
    table.style = "Table Grid"
    if header:
        set_repeat_table_header(table.rows[0])
        for cell in table.rows[0].cells:
            shade_cell(cell, NAVY)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_font(run, 9.2, True, WHITE)
    for row_index, row in enumerate(table.rows):
        set_cant_split_row(row)
        if row_index and row_index % 2 == 0:
            for cell in row.cells:
                shade_cell(cell, "F8FAFC")
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    if row_index:
                        set_font(run, 8.8, False, "24313D")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    set_table_geometry(table, widths)
    style_table(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc: Document, label: str, text: str, fill: str = LIGHT_BLUE,
                label_color: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_table_geometry(table, [9360])
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}  ")
    set_font(r, 10.2, True, label_color)
    r = p.add_run(text)
    set_font(r, 10.2, False, "24313D")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 10.5, True, NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, 10.5, False, "24313D")
    else:
        r = p.add_run(text)
        set_font(r, 10.5, False, "24313D")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1


def add_prompt_box(doc: Document, title: str, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    set_font(r, 11, True, NAVY)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, "F7F9FB")
    set_table_geometry(table, [9360])
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, 8.8, False, "303942", "Microsoft YaHei")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, {1: 16, 2: 13, 3: 11.5}[level], True,
             BLUE if level < 3 else NAVY)


def _font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    chosen = candidates[1] if bold and candidates[1].exists() else next(p for p in candidates if p.exists())
    return ImageFont.truetype(str(chosen), size=size)


def draw_box(draw, xy, text, fill, outline=GRID, title=False, radius=16):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=f"#{outline}", width=2)
    font = _font(24 if title else 20, title)
    x0, y0, x1, y1 = xy
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=6, align="center")
    tx = (x0 + x1 - (bbox[2] - bbox[0])) / 2
    ty = (y0 + y1 - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((tx, ty), text, font=font, fill="#17324D", spacing=6, align="center")


def make_workflow_image(path: Path) -> None:
    img = Image.new("RGB", (1500, 930), "white")
    d = ImageDraw.Draw(img)
    title = _font(30, True)
    d.text((55, 35), "ETF主题雷达：一次研究任务的端到端流程", font=title, fill="#17324D")
    boxes = [
        ((70, 120, 390, 245), "用户提交主题\n与研究目标", "#E8F1F8"),
        ((590, 120, 910, 245), "LLM 规划主题边界\n输出结构化定义", "#E8F1F8"),
        ((1110, 120, 1430, 245), "人工审核主题定义\n批准 / 退回", "#FFF3D9"),
        ((1110, 390, 1430, 545), "研究 Agent 自主选工具\n已有证据 → 官方源 → 公网\n反方检查 → finish", "#EAF5EF"),
        ((590, 390, 910, 545), "确定性治理与评分\n来源分级、证据门槛\n组件评分与惩罚项", "#F2F4F7"),
        ((70, 390, 390, 545), "LLM 证据约束分析\n只使用 evidence_id\n生成研究含义与摘要", "#E8F1F8"),
        ((70, 695, 390, 820), "确定性报告审计\n数字与引用一致性", "#F2F4F7"),
        ((590, 695, 910, 820), "人工审核最终报告\n批准 / 退回重跑", "#FFF3D9"),
        ((1110, 695, 1430, 820), "持久化报告版本\nclaims ↔ evidence", "#EAF5EF"),
    ]
    for xy, text, fill in boxes:
        draw_box(d, xy, text, fill, title=False)
    arrows = [
        ((390, 182), (590, 182)), ((910, 182), (1110, 182)),
        ((1270, 245), (1270, 390)), ((1110, 467), (910, 467)),
        ((590, 467), (390, 467)), ((230, 545), (230, 695)),
        ((390, 757), (590, 757)), ((910, 757), (1110, 757)),
    ]
    for start, end in arrows:
        d.line([start, end], fill="#68717A", width=5)
        ex, ey = end
        sx, sy = start
        if ex > sx:
            pts = [(ex, ey), (ex - 16, ey - 10), (ex - 16, ey + 10)]
        elif ex < sx:
            pts = [(ex, ey), (ex + 16, ey - 10), (ex + 16, ey + 10)]
        else:
            pts = [(ex, ey), (ex - 10, ey - 16), (ex + 10, ey - 16)]
        d.polygon(pts, fill="#68717A")
    note = _font(17)
    d.text((55, 870), "蓝色：LLM参与    灰色：确定性代码    黄色：人工审批    绿色：持久化/交付", font=note, fill="#68717A")
    img.save(path, quality=95)


def make_state_image(path: Path) -> None:
    img = Image.new("RGB", (1500, 730), "white")
    d = ImageDraw.Draw(img)
    d.text((55, 35), "研究运行状态机（主流程）", font=_font(30, True), fill="#17324D")
    labels = [
        ("planning", "主题规划"), ("awaiting_theme_review", "等待主题审核"),
        ("queued / collecting", "排队 / 采集"), ("governing", "证据治理"),
        ("analyzing", "分析与成稿"), ("auditing", "报告审计"),
        ("awaiting_report_review", "等待报告审核"), ("completed", "完成"),
    ]
    positions = [(55, 140), (405, 140), (755, 140), (1105, 140),
                 (1105, 430), (755, 430), (405, 430), (55, 430)]
    for (machine, cn), (x, y) in zip(labels, positions):
        fill = "#FFF3D9" if "awaiting" in machine else "#EAF5EF" if machine == "completed" else "#E8F1F8"
        draw_box(d, (x, y, x + 285, y + 125), f"{machine}\n{cn}", fill)
    arrows = [
        ((340, 202), (405, 202)), ((690, 202), (755, 202)), ((1040, 202), (1105, 202)),
        ((1247, 265), (1247, 430)), ((1105, 492), (1040, 492)), ((755, 492), (690, 492)),
        ((405, 492), (340, 492)),
    ]
    for start, end in arrows:
        d.line([start, end], fill="#68717A", width=5)
        ex, ey = end
        sx, sy = start
        if ex > sx: pts = [(ex, ey), (ex - 14, ey - 9), (ex - 14, ey + 9)]
        elif ex < sx: pts = [(ex, ey), (ex + 14, ey - 9), (ex + 14, ey + 9)]
        else: pts = [(ex, ey), (ex - 9, ey - 14), (ex + 9, ey - 14)]
        d.polygon(pts, fill="#68717A")
    d.text((55, 625), "旁路状态：blocked_configuration（模型配置阻断）、cancelled（取消）、failed（阶段失败）。", font=_font(19), fill="#9B1C1C")
    d.text((55, 665), "阶段级恢复：Worker 通过 SQLite lease 领取任务；每个阶段记录 run_steps，失败最多重试一次。", font=_font(19), fill="#68717A")
    img.save(path, quality=95)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 8.5, False, MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, 8.5, False, MID_GRAY)


def setup_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level, size, before, after, color in (
        (1, 16, 16, 8, BLUE), (2, 13, 12, 6, BLUE), (3, 11.5, 8, 4, NAVY)
    ):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("ETF THEME RADAR  |  TECHNICAL WORKFLOW")
    set_font(r, 8, True, MID_GRAY, "Calibri")
    add_page_number(sec.footer.paragraphs[0])


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_workflow_image(FLOW_IMAGE)
    make_state_image(STATE_IMAGE)
    doc = Document()
    setup_document(doc)

    # Editorial cover: restrained technical report treatment.
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TECHNICAL ARCHITECTURE BRIEF")
    set_font(r, 10, True, BLUE, "Calibri")
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ETF主题雷达 Agent")
    set_font(r, 28, True, NAVY)
    p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("工作流程、模型职责、状态机、数据与记忆机制")
    set_font(r, 15, False, BLUE)
    p.paragraph_format.space_after = Pt(30)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于当前代码库实现梳理")
    set_font(r, 10.5, False, MID_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(date.today().isoformat())
    set_font(r, 10.5, False, MID_GRAY, "Calibri")
    doc.add_page_break()

    heading(doc, "1. 执行摘要")
    add_callout(doc, "核心结论", "当前系统采用“LLM负责规划、工具选择与证据约束表达；确定性代码负责状态、评分、治理、审计和持久化”的混合架构。模型不能自行改变工作流状态、越过人工审批或直接给出投资建议。")
    add_body(doc, "一次完整研究从主题定义开始，经人工审核后由持久 Worker 采集证据，随后进入治理、评分、分析和审计，最终再次等待人工审核。所有关键动作写入 SQLite，界面通过 SSE 读取持久化审计快照；SSE 不可用时回退为普通轮询。")
    add_table(doc, ["层级", "主要职责", "明确边界"], [
        ["大模型层", "主题边界定义、研究工具选择、反方检查、证据含义与摘要生成", "不能控制状态机、不能绕过审批、不能脱离 evidence_id 造事实"],
        ["确定性业务层", "采集编排、来源治理、评分、审计、重试、幂等与取消", "评分参数来自 config/defaults.yaml，不由模型临时改写"],
        ["持久化层", "证据、任务、步骤、调用、快照、实体、报告版本与 claim-evidence", "当前没有向量数据库，也不保存隐藏思维链"],
        ["人工审核层", "主题定义审批与最终报告审批", "退回后按路由重新进入规划或排队执行"],
    ], [1700, 4300, 3360])

    heading(doc, "2. 端到端工作流程")
    doc.add_picture(str(FLOW_IMAGE), width=Inches(6.7))
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_body(doc, "流程中的两次人工门禁分别位于主题定义之后和报告审计之后。研究 Agent 只能在 collecting 阶段自主选择已注册工具；后续治理、评分、审计均由确定性代码执行。")

    heading(doc, "3. 大模型参与的步骤")
    add_table(doc, ["步骤", "模型 / 接口", "输入", "输出", "确定性约束"], [
        ["主题定义", "DeepSeek V4 Pro（PydanticAI）", "用户主题、研究目标", "description、aliases、include_terms、research_questions", "结构化校验；失败时使用确定性后备定义"],
        ["证据采集规划", "DeepSeek V4 Flash（PydanticAI Agent）", "已审主题定义、已有证据、来源健康、预算", "工具调用序列与停止原因", "工具白名单、预算、幂等、反方检查与 finish_research 防护栏"],
        ["证据含义分析", "DeepSeek V4 Pro（JSON Chat Completions）", "通过治理的 evidence_id 与事实字段", "implications、conflicts、limitations、missing_evidence", "temperature=0；禁止新增 URL、公司或数字"],
        ["报告摘要", "DeepSeek V4 Pro（JSON Chat Completions）", "确定性评分、证据、审计上下文", "受证据约束的摘要文本", "数字与引用由后续审计核对；失败可降级"],
        ["事件深挖", "分析模型 + 确定性核验", "单一 evidence_id 及其来源", "事件事实包与深挖报告", "独立于主题主流程，状态为 source_verification → analyzing → completed"],
    ], [1200, 1700, 1900, 2100, 2460])
    add_callout(doc, "模型配置", "默认环境变量为 LLM_AGENT_MODEL=deepseek-v4-flash、LLM_ANALYSIS_MODEL=deepseek-v4-pro。模型未配置时，系统记录 tools_unavailable；配置错误可进入 blocked_configuration。", LIGHT_GRAY, NAVY)

    heading(doc, "4. 当前提示词")
    add_body(doc, "以下为当前三类核心提示词的内容要点。主题定义与研究 Agent 从 prompts/ 文件加载；证据分析文件存在，但实际分析运行仍使用 theme_research.py 中的内联 JSON 提示词，这是当前需要统一的一处实现差异。")
    add_prompt_box(doc, "4.1 主题定义提示词（theme_definition.md）",
                   "你是 ETF 主题研究规划器。根据用户给出的主题和研究目标，输出一个可执行的研究定义。只定义研究边界，不给出投资、证券、产品或交易建议；aliases 和 include_terms 各保留 1–8 项；research_questions 必须覆盖支持证据、反方证据和可投资性；不生成未由输入提供的 URL、市场规模、收益率或其他数字。")
    add_prompt_box(doc, "4.2 研究 Agent 提示词（research_agent.md）",
                   "你是证据优先的 ETF 主题研究 Agent。先检查已有证据和来源健康；优先官方或学术连接器，再使用公开搜索，浏览器仅读取明确的公开候选页面；每次调用后评估新增证据、独立来源类型和缺口；不得重复无新增价值的调用；至少执行一次反方证据或否定条件搜索；工具失败时切换可用来源且不得虚构结果；达到门槛或预算不足时调用 finish_research。summarize_evidence_gap 后，若尚未反方检查则下一步必须 counter，否则必须 finish_research。")
    add_prompt_box(doc, "4.3 证据分析提示词（evidence_analysis.md / 运行时内联版本）",
                   "你是受证据约束的 ETF 主题研究分析器。每个事实性判断必须绑定已有 evidence_id；不得新增输入中不存在的数字、公司、URL 或事实；明确支持证据、冲突、限制和仍缺失的证据；不提供投资、产品、组合或交易建议。")

    heading(doc, "5. 状态机与流转规则")
    doc.add_picture(str(STATE_IMAGE), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(doc, ["当前状态", "触发 / 执行", "下一状态", "持久化动作"], [
        ["planning", "生成主题定义", "awaiting_theme_review", "写入 planning run_step、主题定义和 review_gate"],
        ["awaiting_theme_review", "人工批准", "queued", "CAS 原子状态更新并写 approvals"],
        ["queued / collecting", "研究 Agent 调用工具采集", "governing", "写 agent_runs、tool_calls、raw_documents、normalized_events"],
        ["governing", "来源分级、门槛与证据治理", "analyzing", "更新事件治理字段和阶段结果"],
        ["analyzing", "确定性评分 + LLM 证据分析", "auditing", "写主题分数、快照、报告候选与引用"],
        ["auditing", "数字、引用与结构审计", "awaiting_report_review", "写 report-audit；即使失败也进入人工审核并标记 audit_passed=false"],
        ["awaiting_report_review", "人工批准", "completed", "注册 report_assets、report_versions、claims 与 claim-evidence"],
        ["任一可取消状态", "用户取消 / 竞态 CAS", "cancelled", "保留已完成步骤与审计记录"],
    ], [1750, 2350, 1900, 3360])

    heading(doc, "6. Worker、恢复与幂等")
    add_body(doc, "Worker 模式：单进程持久 Worker 从 SQLite 领取任务；API 只负责创建任务或执行原子状态操作，不为每个请求创建线程。")
    add_body(doc, "租约机制：任务通过 SQLite lease 认领，默认租约约 30 秒，心跳间隔约为租约的三分之一。进程异常后，过期租约可被重新领取。")
    add_body(doc, "阶段恢复：run_steps 以 run_id、attempt、step_name 为主键记录阶段。每次执行只推进一个阶段并释放租约；同一阶段异常时最多重试一次，再转为 failed。")
    add_body(doc, "调用幂等：稳定外部工具调用以 run_id、attempt、工具名和规范化参数生成 call_uid；重复调用优先重放已持久化结果。失败、跳过和拒绝也写入 tool_calls。")

    heading(doc, "7. 数据存储与目录")
    add_callout(doc, "默认数据库", "data/radar.db（SQLite）。它是任务恢复、审计、证据引用和报告版本的事实来源。", LIGHT_BLUE, BLUE)
    add_table(doc, ["数据域", "主要表", "保存内容"], [
        ["原始证据与事件", "raw_documents, normalized_events, citations", "原文、URL、哈希、抓取/发布时间、解析版本、来源类型、主题/公司/代码、质量与置信度"],
        ["来源运行状态", "connector_health, sync_runs", "连接器启用状态、覆盖说明、同步进度、错误与取消请求"],
        ["主题与趋势", "themes, theme_aliases, theme_scores, theme_snapshots", "主题定义、审核别名、评分分量、证据集合、至少两个快照后才可判断趋势"],
        ["实体关系", "entities, entity_aliases, entity_links", "公司/ETF 等实体、别名、ticker、exchange、实体—事件—主题关系与复核状态"],
        ["运行与审计", "research_runs, run_steps, approvals, agent_runs, tool_calls, run_registry", "状态、阶段、尝试次数、lease、模型与 token 使用、工具参数/结果/耗时/重试/证据增量"],
        ["报告与可追溯性", "report_assets, report_versions, report_claims, report_claim_evidence", "报告资产、版本正文、内容哈希、事实 claim 及其 evidence_id 关系"],
        ["产品研究候选", "product_proposals", "主题对应的结构化候选载荷；不等同投资或交易建议"],
    ], [1750, 3000, 4610])
    add_body(doc, "文件制品位于 data/reports/research-runs/<run_id>/attempt-<n>/，可包括 theme-hypothesis.json、theme-score.json、investability-assessment.json、etf-landscape.json、report-audit.json、evidence-appendix.json 与 Markdown 报告；事件深挖另保存 event-fact-pack.json 和 event-deep-dive.md。")

    heading(doc, "8. 研究工具与 Skill")
    add_table(doc, ["类别", "工具 / Skill", "用途", "当前调用状态"], [
        ["Agent 工具", "query_existing_evidence", "读取已有证据与覆盖情况", "运行时可调用；当前返回全局前若干事件，主题过滤仍需加强"],
        ["Agent 工具", "inspect_source_health", "检查已注册连接器健康与可用性", "运行时可调用"],
        ["Agent 工具", "collect_from_source", "调用 sec、arxiv、patents、etf_holdings、company_careers、sp_global", "运行时可调用；各连接器按配置启停"],
        ["Agent 工具", "search_public_web", "公开网页发现与反方检索", "通过 AnySearch 补充；只作为发现层"],
        ["Agent 工具", "browse_public_page", "读取明确的公开候选页面", "默认关闭；Playwright 正文按不可信输入清洗"],
        ["Agent 工具", "summarize_evidence_gap", "汇总缺口并约束下一动作", "运行时可调用"],
        ["Agent 工具", "finish_research", "记录停止原因和剩余缺口", "最终结构化输出前强制调用"],
        ["项目 Skill", "anysearch", "实时网页、垂直搜索和公开页面提取", "当前 Worker 实际调用的项目 Skill"],
        ["已审计 Skill", "etf-premium / stock-liquidity / yfinance-data", "溢折价、流动性与公开市场数据框架", "尚未接入 Worker；相关判断保持 not_assessed"],
        ["已审计 Skill", "idea-generation / competitive-analysis / sector-overview", "主题构思、竞品与行业研究框架", "作为人工研究参考，尚未由 Worker 自动调用"],
    ], [1300, 2450, 3300, 2310])
    add_callout(doc, "概念区分", "Connector 是项目中的数据源适配器；Tool 是模型可调用的受审计函数；Skill 是可复用工作方法或外部能力包。三者不是同一层。", LIGHT_GRAY, NAVY)

    heading(doc, "9. 预算、防护栏与停止条件")
    add_table(doc, ["项目", "常规研究默认值", "Quick Scan", "说明"], [
        ["最大工具调用", "16", "最多 6", "连续两次无新增证据后停止重复支持性采集"],
        ["最大模型请求", "8", "最多 4", "由 PydanticAI UsageLimits 与本地预算共同限制"],
        ["最大运行时间", "标准档 480 秒", "ETF 快速档 240 秒", "同步外部调用完成后才能响应取消"],
        ["最大 token", "24,000", "24,000 上限内受请求数限制", "同时受成本上限折算的 token 限制"],
        ["最大估算成本", "0.50 美元", "同一上限", "估算费率默认为 2 美元 / 百万 token"],
    ], [1800, 1700, 1800, 4060])
    add_body(doc, "允许的停止原因包括 evidence_sufficient、insufficient_evidence、budget_exhausted、tools_unavailable。若未完成 counter 检查，finish_research 会被拒绝；summarize_evidence_gap 之后只允许一次明确 counter 或直接 finish。")

    heading(doc, "10. 记忆机制：系统记得什么，不记得什么")
    add_table(doc, ["记忆范围", "载体", "具体内容", "生命周期"], [
        ["单次模型会话", "PydanticAI / conversation_id", "同一 agent_run 内的工具调用上下文与结构化输出上下文", "仅当前 agent_run；不作为跨任务长期记忆"],
        ["任务记忆", "SQLite research_runs / run_steps / tool_calls", "状态、阶段、尝试、参数、结果、错误、停止原因与预算使用", "持久保存，可用于恢复和审计"],
        ["领域记忆", "themes / aliases / entities / snapshots", "审核后的主题边界、实体规范名、别名、关系和可比历史快照", "跨运行持久保存"],
        ["证据记忆", "raw_documents / normalized_events / citations", "去重原文、来源元数据、规范化事件及证据质量", "跨运行持久保存"],
        ["报告记忆", "report_versions / claims / claim_evidence", "报告版本、内容哈希、事实主张与证据映射", "跨运行持久保存"],
        ["不存在的记忆", "无", "无向量数据库、无 ChatGPT/Codex 私人记忆调用、无跨任务隐藏对话历史、无隐藏思维链存储", "不保存"],
    ], [1450, 2300, 3700, 1910])

    heading(doc, "11. 当前已知缺口与优先级建议")
    add_table(doc, ["优先级", "缺口", "影响", "建议"], [
        ["P1", "query_existing_evidence 返回全局前若干事件，而非严格按主题筛选", "模型可能看到低相关证据；覆盖统计与返回内容不完全一致", "按 theme_id、aliases、include_terms 做 SQL/检索层过滤并增加契约测试"],
        ["P1", "提示词文件与 theme_research.py 内联提示词并存", "版本审计与修改容易漂移", "统一由 prompts/ 加载，并把 prompt_hash 写入所有分析调用"],
        ["P1", "审计失败仍进入最终人工审核", "流程可见但容易被误认为已通过", "前端强提示 audit_passed=false；按风险决定是否增设 blocked_audit"],
        ["P2", "没有向量检索或语义召回", "证据量扩大后，关键词与前若干条读取会降低召回质量", "先实现主题过滤与 FTS，再评估是否需要 embeddings"],
        ["P2", "摘要仅校验 evidence_id 与数字，不做完整语义蕴含验证", "引用存在不代表 claim 被证据真正支持", "增加 claim-evidence entailment 校验器与人工抽样规则"],
        ["P2", "取消只能在同步外部调用返回后生效", "慢请求期间取消延迟", "连接器采用可中断超时、分段请求和阶段间取消检查"],
        ["P3", "流动性、溢折价、指数规则相关 Skills 尚未接入 Worker", "ETF 可投资性部分字段只能保持 unknown / not_assessed", "在正式数据授权与来源治理明确后，通过可替换 adapter 接入"],
    ], [900, 3000, 2800, 2660])

    doc.add_page_break()
    heading(doc, "12. 关键实现位置")
    add_table(doc, ["主题", "文件"], [
        ["工作流与状态转换", "etf_theme_radar/workflow.py"],
        ["持久 Worker 与 lease", "etf_theme_radar/worker.py"],
        ["研究 Agent、工具、防护栏与预算", "etf_theme_radar/agent_runtime.py"],
        ["SQLite schema 与数据访问", "etf_theme_radar/store.py"],
        ["证据分析与报告生成", "etf_theme_radar/theme_research.py"],
        ["确定性评分", "etf_theme_radar/scoring.py；config/defaults.yaml"],
        ["提示词", "etf_theme_radar/prompts/*.md"],
        ["Skill 审计", "docs/skill-audit.md"],
        ["API、SSE 与人工审核路由", "etf_theme_radar/api.py"],
    ], [3000, 6360])
    add_callout(doc, "文档边界", "本说明描述的是当前代码库的真实实现与已识别缺口，不代表未来规划已经落地；所有 ETF 持仓、流动性、指数规则及美国可交易状态在未核验时仍应显示 unknown / not_assessed。", "FFF3D9", AMBER)

    doc.core_properties.title = "ETF主题雷达 Agent 工作流程与技术架构说明"
    doc.core_properties.subject = "LLM职责、状态机、数据存储、工具与记忆机制"
    doc.core_properties.author = "ETF Theme Radar"
    doc.core_properties.keywords = "ETF, Agent, Workflow, State Machine, SQLite, LLM"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
